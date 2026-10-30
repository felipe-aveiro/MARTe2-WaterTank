import pandas as pd
from docx import Document
from docx.shared import Pt
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ===== CONFIGURATION =====
INPUT_CSV = "C:\\Users\\felip\\Desktop\\MARTe2-WaterTank\\DataVisualization\\Outputs\\PlasmaPositionMetrics\\plasma_position_metrics_all_selected_shots.csv"
OUTPUT_DOCX = "metrics_table.docx"

shot_order = [45967, 46241, 48555, 48556,
              48561, 48563, 53058, 53071]

position_order = ["Radial", "Vertical"]
estimator_order = ["Magnetic", "Electric", "Fused"]

df = pd.read_csv(INPUT_CSV)
df.columns = df.columns.str.strip()

df["position"] = df["position"].str.strip()
df["position"] = df["position"].replace({
    "R": "Radial",
    "Z": "Vertical"
})
df["estimator"] = df["estimator"].str.strip()

# ===== VALIDATION =====
invalid = df[
    (df["MAE_cm"] > df["RMSE_cm"] + 1e-10) |
    (df["RMSE_cm"] > df["MaxAE_cm"] + 1e-10) |
    (df["PCC"].abs() > 1 + 1e-10)
]

if not invalid.empty:
    print("ATTENTION: inconsistencies found:")
    print(invalid.to_string(index=False))
    raise ValueError("Inconsistent metrics in CSV.")

print("All metrics validated.")

df = df[df["shot"].isin(shot_order)].copy()

df["shot"] = pd.Categorical(
    df["shot"], categories=shot_order, ordered=True
)
df["position"] = pd.Categorical(
    df["position"], categories=position_order, ordered=True
)
df["estimator"] = pd.Categorical(
    df["estimator"], categories=estimator_order, ordered=True
)

df = df.sort_values(
    ["shot", "position", "estimator"]
).reset_index(drop=True)

# ===== WORD =====
doc = Document()

table = doc.add_table(rows=1, cols=8)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = [
    "Shot number", "Analyzed interval [ms]",
    "Position", "Estimator", "MAE [cm]",
    "RMSE [cm]", "MaxAE [cm]", "PCC"
]

for cell, text in zip(table.rows[0].cells, headers):
    cell.text = text

for _, r in df.iterrows():
    cells = table.add_row().cells

    interval = (
        f"{int(r['flat_top_start_ms'])} - "
        f"{int(r['flat_top_end_ms'])}"
    )

    values = [
        str(int(r["shot"])),
        interval,
        str(r["position"]),
        str(r["estimator"]),
        f"{r['MAE_cm']:.3f}",
        f"{r['RMSE_cm']:.3f}",
        f"{r['MaxAE_cm']:.3f}",
        f"{r['PCC']:.3f}"
    ]

    for cell, value in zip(cells, values):
        cell.text = value

for shot in shot_order:
    indexes = df.index[df["shot"] == shot].tolist()

    if not indexes:
        continue

    first = indexes[0] + 1
    last = indexes[-1] + 1

    for col in [0, 1]:
        table.cell(first, col).merge(table.cell(last, col))

    # (Radial / Vertical)
    for pos in position_order:
        pos_indexes = df.index[
            (df["shot"] == shot) &
            (df["position"] == pos)
        ].tolist()

        if pos_indexes:
            a = pos_indexes[0] + 1
            b = pos_indexes[-1] + 1
            table.cell(a, 2).merge(table.cell(b, 2))

# ===== FORMATTING =====
for row in table.rows:
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(9)

for cell in table.rows[0].cells:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True

doc.save(OUTPUT_DOCX)

print(f"Table created: {OUTPUT_DOCX}")
